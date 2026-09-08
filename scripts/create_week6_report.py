from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# PATHS
# ============================================================

ROOT = Path(__file__).resolve().parent.parent

SCREENSHOTS = ROOT / "screenshots" / "week6"
OUTPUT = ROOT / "reports" / "Week6_Integrative_Capstone_Report.docx"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)


# ============================================================
# DOCUMENT SETUP
# ============================================================

doc = Document()

section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles

styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)

styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(22)

styles["Heading 1"].font.name = "Arial"
styles["Heading 1"].font.size = Pt(16)

styles["Heading 2"].font.name = "Arial"
styles["Heading 2"].font.size = Pt(13)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(22)

    return p


def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(12)

    return p


def add_heading(text, level=1):
    doc.add_heading(text, level=level)


def add_para(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_number(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_image(filename, caption, width=6.2):
    path = SCREENSHOTS / filename

    if not path.exists():
        print(f"[WARNING] Image not found: {path}")
        add_para(f"[Figure not found: {filename}]")
        return False

    print(f"[OK] Embedding image: {path}")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)

    return True


def add_page_break():
    doc.add_page_break()


def add_table(headers, rows):
    table = doc.add_table(
        rows=1,
        cols=len(headers)
    )

    table.style = "Table Grid"

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)

    for row in rows:
        cells = table.add_row().cells

        for i, value in enumerate(row):
            cells[i].text = str(value)

    return table


# ============================================================
# TITLE PAGE
# ============================================================

add_title(
    "Adult Income Prediction and Data Science Analysis"
)

add_subtitle(
    "Week 6 — Integrative Capstone Project and Evaluation"
)

doc.add_paragraph()
add_subtitle(
    "End-to-End Python Data Science, Machine Learning and Deep Learning Pipeline"
)

doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

r = p.add_run(
    "Internship Capstone Report"
)
r.bold = True
r.font.size = Pt(14)

doc.add_page_break()


# ============================================================
# 1. EXECUTIVE SUMMARY
# ============================================================

add_heading("1. Executive Summary")

add_para(
    "This final capstone integrates the complete data science workflow "
    "developed throughout the six-week internship using the Adult Income "
    "dataset. The project combines data preparation, exploratory data "
    "analysis, unsupervised learning, supervised machine learning, and "
    "deep learning into a single analytical pipeline."
)

add_para(
    "The primary objective is to understand factors associated with income "
    "classification and predict whether an individual belongs to the "
    "`<=50K` or `>50K` annual income category."
)

add_para(
    "The project evaluates multiple traditional machine-learning models "
    "along with a neural network. The final neural network achieved "
    "83.06% accuracy, 60.49% precision, 85.54% recall, 70.87% F1-score, "
    "and 91.90% ROC-AUC."
)

add_para(
    "The analysis demonstrates that model selection depends on the "
    "application objective. Random Forest provided stronger overall "
    "accuracy and precision, while the neural network provided strong "
    "recall for the minority `>50K` class."
)


# ============================================================
# 2. PROBLEM STATEMENT
# ============================================================

add_heading("2. Problem Statement")

add_para(
    "Income classification is a useful example of a binary classification "
    "problem involving demographic, educational, employment, and financial "
    "information. The goal of this project is to develop a complete data "
    "science workflow capable of analyzing the factors associated with "
    "income and predicting whether an individual's annual income is "
    "greater than $50K."
)

add_para(
    "The problem is challenging because the target classes are imbalanced. "
    "There are substantially more observations in the `<=50K` category "
    "than in the `>50K` category. Therefore, model performance must be "
    "evaluated using multiple metrics rather than accuracy alone."
)


# ============================================================
# 3. DATASET
# ============================================================

add_heading("3. Dataset and Data Preparation")

add_para(
    "The project uses the cleaned Adult Income dataset developed during "
    "the earlier stages of the internship. The final dataset contains "
    "32,537 observations and 15 original variables, with `income` serving "
    "as the target variable."
)

add_para(
    "The dataset contains demographic variables such as age, sex, race, "
    "and relationship; educational variables such as education and "
    "education number; employment variables such as workclass, occupation, "
    "and hours per week; and financial variables including capital gain "
    "and capital loss."
)

add_heading("Target Distribution", level=2)

add_para(
    "The target distribution contains 24,698 individuals in the `<=50K` "
    "category and 7,839 individuals in the `>50K` category. This corresponds "
    "to approximately 75.91% and 24.09%, respectively."
)

add_image(
    "week6_target_distribution.png",
    "Figure 1. Distribution of the Adult Income target classes."
)


# ============================================================
# 4. EDA
# ============================================================

add_heading("4. Exploratory Data Analysis")

add_para(
    "Exploratory data analysis was performed to identify patterns and "
    "relationships that could help explain income classification before "
    "building predictive models."
)

add_heading("4.1 Education and Income", level=2)

add_para(
    "Educational attainment showed a meaningful association with income. "
    "Education levels associated with higher qualifications generally "
    "contained a larger proportion of individuals earning more than $50K."
)

add_image(
    "week6_education_income.png",
    "Figure 2. Income distribution across education levels."
)

add_heading("4.2 Age and Income", level=2)

add_para(
    "The higher-income group generally showed an older age distribution. "
    "This pattern may reflect differences in work experience, career "
    "progression, and earning potential."
)

add_image(
    "week6_age_income.png",
    "Figure 3. Age distribution by income category."
)

add_heading("4.3 Working Hours and Income", level=2)

add_para(
    "Individuals in the higher-income category generally showed higher "
    "weekly working hours. Working hours therefore provide useful "
    "information for income classification, although they cannot "
    "independently determine income."
)

add_image(
    "week6_hours_income.png",
    "Figure 4. Weekly working hours by income category."
)


# ============================================================
# 5. UNSUPERVISED LEARNING
# ============================================================

add_heading("5. Unsupervised Learning — K-Means and PCA")

add_para(
    "K-Means clustering was used to identify naturally occurring groups "
    "within the dataset. The clustering stage used age, education number, "
    "hours per week, capital gain, and capital loss."
)

add_para(
    "The features were standardized before clustering because they operate "
    "on different numerical scales. Silhouette-score analysis supported "
    "the selection of K=2."
)

add_para(
    "The resulting clusters showed meaningful differences. One smaller "
    "cluster contained approximately 1,484 observations, while the larger "
    "cluster contained approximately 31,053 observations."
)

add_para(
    "The smaller cluster had an average age of approximately 41.63 years, "
    "education number of 10.99, and 43.36 working hours per week. "
    "Approximately 51.95% of this group belonged to the `>50K` category."
)

add_para(
    "The larger cluster had an average age of approximately 38.44 years, "
    "education number of 10.04, and 40.30 working hours per week. "
    "Approximately 22.76% belonged to the `>50K` category."
)

add_para(
    "This represents a difference of approximately 29.19 percentage "
    "points in the proportion of higher-income individuals between "
    "the two clusters."
)

add_image(
    "week6_cluster_pca.png",
    "Figure 5. PCA visualization of the K-Means clusters."
)


# ============================================================
# 6. SUPERVISED LEARNING
# ============================================================

add_heading("6. Supervised Learning Model Comparison")

add_para(
    "Multiple supervised-learning algorithms were evaluated during "
    "the project. Logistic Regression, Decision Tree, Random Forest, "
    "and Balanced Random Forest were compared with the final neural "
    "network."
)

headers = [
    "Model",
    "Accuracy",
    "Precision",
    "Recall",
    "F1",
    "ROC-AUC"
]

rows = [
    ["Logistic Regression", "85.79%", "74.19%", "62.88%", "68.07%", "91.47%"],
    ["Decision Tree", "86.60%", "78.81%", "60.71%", "68.59%", "90.60%"],
    ["Random Forest", "87.11%", "82.17%", "59.38%", "68.94%", "92.30%"],
    ["Balanced Random Forest", "81.78%", "58.08%", "87.56%", "69.84%", "92.33%"],
    ["Neural Network", "83.06%", "60.49%", "85.54%", "70.87%", "91.90%"],
]

add_table(headers, rows)

doc.add_paragraph()

add_image(
    "week6_model_comparison.png",
    "Figure 6. Comparison of predictive-model performance."
)

add_para(
    "Random Forest achieved the highest accuracy at 87.11% and the "
    "highest precision at 82.17%. Balanced Random Forest achieved the "
    "highest recall among the traditional and neural models at 87.56% "
    "and the highest ROC-AUC at 92.33%."
)

add_para(
    "The neural network achieved the highest F1-score in the comparison "
    "at 70.87%, while maintaining a strong recall of 85.54%."
)


# ============================================================
# 7. DEEP LEARNING
# ============================================================

add_heading("7. Deep Learning Model")

add_para(
    "The final deep-learning model is a feed-forward neural network "
    "implemented using TensorFlow and Keras."
)

add_heading("Architecture", level=2)

add_para(
    "The network contains an input layer followed by two fully connected "
    "hidden layers. The first hidden layer contains 64 neurons with ReLU "
    "activation, followed by 30% dropout. The second hidden layer contains "
    "32 neurons with ReLU activation, followed by 20% dropout. A single "
    "sigmoid output neuron performs the binary classification."
)

add_para(
    "Dropout was included to reduce overfitting. The Adam optimizer was "
    "used with a learning rate of 0.001, while binary cross-entropy was "
    "used as the loss function."
)

add_para(
    "Class weights were applied because the `>50K` class is the minority "
    "class. Early stopping was also used to prevent unnecessary training "
    "after validation performance stopped improving."
)


# ============================================================
# 8. FINAL EVALUATION
# ============================================================

add_heading("8. Final Neural Network Evaluation")

add_para(
    "The final Week 6 neural network was evaluated on 4,881 previously "
    "unseen test observations."
)

headers = ["Metric", "Result"]

rows = [
    ["Accuracy", "83.06%"],
    ["Precision", "60.49%"],
    ["Recall", "85.54%"],
    ["F1-score", "70.87%"],
    ["ROC-AUC", "91.90%"],
]

add_table(headers, rows)

doc.add_paragraph()

add_heading("8.1 Confusion Matrix", level=2)

add_para(
    "The confusion matrix provides a detailed breakdown of the model's "
    "predictions."
)

headers = [
    "",
    "Predicted <=50K",
    "Predicted >50K"
]

rows = [
    ["Actual <=50K", "3,048", "657"],
    ["Actual >50K", "170", "1,006"],
]

add_table(headers, rows)

doc.add_paragraph()

add_para(
    "The model correctly classified 3,048 observations in the <=50K "
    "category and 1,006 observations in the >50K category. It produced "
    "657 false positives and 170 false negatives."
)

add_image(
    "week6_confusion_matrix.png",
    "Figure 7. Confusion matrix for the final neural network."
)

add_heading("8.2 ROC Curve", level=2)

add_para(
    "The final neural network achieved a ROC-AUC of 91.90%, indicating "
    "strong discrimination between the two income categories across "
    "different classification thresholds."
)

add_image(
    "week6_roc_curve.png",
    "Figure 8. ROC curve for the final neural network."
)


# ============================================================
# 9. ERROR ANALYSIS
# ============================================================

add_heading("9. Error Analysis")

add_para(
    "Error analysis was performed to understand the characteristics "
    "of correctly and incorrectly classified observations."
)

add_para(
    "The analysis compares numerical characteristics including age, "
    "education number, working hours, capital gain, capital loss, "
    "total capital, and the work-hours-to-age ratio."
)

add_image(
    "week6_error_analysis.png",
    "Figure 9. Comparison of feature characteristics for correct and incorrect predictions."
)

add_para(
    "The results indicate that classification errors can occur when "
    "individuals have characteristics that overlap between the two "
    "income categories. This demonstrates why income classification "
    "cannot be reliably determined using a single variable."
)


# ============================================================
# 10. KEY FINDINGS
# ============================================================

add_heading("10. Key Findings")

add_bullet(
    "The dataset contains 32,537 observations with a substantial "
    "imbalance between the two income classes."
)

add_bullet(
    "Education, age, working hours, and financial variables provide "
    "useful information for distinguishing income categories."
)

add_bullet(
    "K-Means clustering identified two groups with substantially "
    "different characteristics and higher-income proportions."
)

add_bullet(
    "Random Forest achieved the highest overall accuracy at 87.11%."
)

add_bullet(
    "Balanced Random Forest achieved the highest recall among the "
    "traditional models at 87.56%."
)

add_bullet(
    "The final neural network achieved 83.06% accuracy, 60.49% "
    "precision, 85.54% recall, 70.87% F1-score, and 91.90% ROC-AUC."
)

add_bullet(
    "The neural network achieved strong minority-class detection "
    "but did not outperform Random Forest in overall accuracy."
)

add_bullet(
    "Model selection should depend on whether the application "
    "prioritizes accuracy, precision, recall, or balanced performance."
)


# ============================================================
# 11. RECOMMENDATIONS
# ============================================================

add_heading("11. Recommendations")

add_number(
    "Use Random Forest when overall prediction accuracy and precision "
    "are the primary objectives."
)

add_number(
    "Consider Balanced Random Forest or the neural network when "
    "identifying as many >50K individuals as possible is more important."
)

add_number(
    "Evaluate models using precision, recall, F1-score, ROC-AUC, and "
    "confusion matrices rather than accuracy alone."
)

add_number(
    "Retain education, age, employment, and financial variables because "
    "income classification depends on multiple characteristics."
)

add_number(
    "Tune the neural-network classification threshold in future work "
    "if a different precision-recall trade-off is required."
)

add_number(
    "Use model-explainability methods such as SHAP in future versions "
    "to improve interpretation of individual predictions."
)


# ============================================================
# 12. LIMITATIONS
# ============================================================

add_heading("12. Limitations")

add_bullet(
    "The Adult Income dataset is a historical benchmark dataset and "
    "may not represent current economic conditions."
)

add_bullet(
    "The target classes are imbalanced, which makes minority-class "
    "prediction more challenging."
)

add_bullet(
    "The neural network does not automatically outperform traditional "
    "ensemble methods on this tabular dataset."
)

add_bullet(
    "Deep-learning experiments were performed in a local development "
    "environment with a relatively compact architecture."
)

add_bullet(
    "Neural-network predictions are less directly interpretable than "
    "simpler statistical models."
)

add_bullet(
    "The evaluation is based on a single held-out test set and would "
    "benefit from external validation in a real-world application."
)


# ============================================================
# 13. FUTURE WORK
# ============================================================

add_heading("13. Future Improvements")

add_bullet(
    "Perform systematic hyperparameter optimization for the neural network."
)

add_bullet(
    "Experiment with additional architectures and regularization methods."
)

add_bullet(
    "Investigate threshold optimization to improve the precision-recall trade-off."
)

add_bullet(
    "Apply SHAP or other explainability methods to understand individual predictions."
)

add_bullet(
    "Evaluate the models using external or more recent income datasets."
)

add_bullet(
    "Develop a deployment interface where users can enter feature values "
    "and receive an income-category prediction."
)


# ============================================================
# 14. FINAL CONCLUSION
# ============================================================

add_heading("14. Final Conclusion")

add_para(
    "This six-week internship project successfully developed an end-to-end "
    "data science pipeline for analyzing and predicting Adult Income categories."
)

add_para(
    "The project progressed from data cleaning and exploratory analysis "
    "to unsupervised learning, supervised machine learning, and deep learning. "
    "This progression provided practical experience across the major stages "
    "of a modern data science workflow."
)

add_para(
    "The exploratory analysis identified meaningful relationships between "
    "income and variables such as education, age, working hours, and financial "
    "characteristics. K-Means clustering provided an additional perspective "
    "by identifying population groups with different income distributions."
)

add_para(
    "The supervised-learning comparison demonstrated that Random Forest "
    "provided the strongest overall accuracy and precision, while Balanced "
    "Random Forest and the neural network provided stronger recall for the "
    "minority income class."
)

add_para(
    "The final neural network achieved 83.06% accuracy, 60.49% precision, "
    "85.54% recall, 70.87% F1-score, and 91.90% ROC-AUC. Its performance "
    "demonstrates that deep learning can effectively learn useful patterns "
    "from tabular data, although a traditional ensemble model remained "
    "stronger for overall accuracy."
)

add_para(
    "Overall, the capstone demonstrates the ability to independently design "
    "and execute a complete Python-based data science project, evaluate "
    "alternative analytical approaches, interpret model behavior, and "
    "translate quantitative results into practical recommendations."
)


# ============================================================
# 15. PROJECT DELIVERABLES
# ============================================================

add_heading("15. Project Deliverables")

add_bullet(
    "Week 1 — Data exploration and cleaning"
)

add_bullet(
    "Week 2 — Exploratory data analysis and visualization"
)

add_bullet(
    "Week 3 — K-Means clustering and PCA"
)

add_bullet(
    "Week 4 — Supervised machine learning"
)

add_bullet(
    "Week 5 — Deep learning application"
)

add_bullet(
    "Week 6 — Integrative capstone and evaluation"
)

add_para(
    "GitHub Repository: https://github.com/bhargav55678/adult-income-data-cleaning"
)


# ============================================================
# SAVE
# ============================================================

doc.save(OUTPUT)

print()
print("=" * 60)
print("WEEK 6 REPORT CREATED SUCCESSFULLY")
print("=" * 60)
print(f"Report: {OUTPUT}")
print(f"Report exists: {OUTPUT.exists()}")
print(f"Report size: {OUTPUT.stat().st_size:,} bytes")
print("=" * 60)